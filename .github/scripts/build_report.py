"""Turn one run's Allure results into the CI report: run summary, email body and Excel sheet.

Usage: python3 .github/scripts/build_report.py [allure-results-dir] [output-dir]

Writes to output-dir (default target/):
  OrangeHRM-API-Test-Results.xlsx  - Summary sheet + one row per test case
  email-body.html                  - counts, failed tests with their errors
Also appends a markdown summary to $GITHUB_STEP_SUMMARY and the counts to $GITHUB_OUTPUT
when those are set (i.e. inside GitHub Actions).
"""
import glob
import html
import json
import os
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

RESULTS_DIR = sys.argv[1] if len(sys.argv) > 1 else "target/allure-results"
OUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "target"
SHEET_NAME = "OrangeHRM-API-Test-Results.xlsx"

# Allure statuses: failed = assertion failure, broken = unexpected exception
STATUS_LABEL = {"passed": "Passed", "failed": "Failed", "broken": "Error", "skipped": "Skipped"}
STATUS_FILL = {"Passed": "C6EFCE", "Failed": "FFC7CE", "Error": "FFC7CE", "Skipped": "FFEB9C"}


def load_results():
    """Latest result per test method (TestNG retries / repeated local runs write several)."""
    latest = {}
    for path in glob.glob(os.path.join(RESULTS_DIR, "*-result.json")):
        with open(path, encoding="utf-8") as f:
            r = json.load(f)
        key = r.get("fullName") or r.get("name")
        if key not in latest or r.get("stop", 0) > latest[key].get("stop", 0):
            latest[key] = r

    rows = []
    for r in latest.values():
        name = r.get("name", "")
        tc_id, _, scenario = name.partition(" | ")
        if not scenario:
            tc_id, scenario = "", name
        cls = (r.get("fullName") or "").rsplit(".", 2)[-2] if r.get("fullName") else ""
        details = r.get("statusDetails") or {}
        rows.append({
            "tc_id": tc_id.strip(),
            "module": cls[:-4] if cls.endswith("Test") else cls,
            "scenario": scenario.strip(),
            "method": (r.get("fullName") or "").rsplit(".", 1)[-1],
            "status": STATUS_LABEL.get(r.get("status"), r.get("status", "Unknown")),
            "duration": round((r.get("stop", 0) - r.get("start", 0)) / 1000, 2),
            "error": (details.get("message") or "").strip(),
            "start": r.get("start", 0),
        })
    rows.sort(key=lambda x: (x["status"] == "Passed", x["module"], x["tc_id"]))
    return rows


def first_lines(text, limit=400):
    """The assertion headline, without the multi-kilobyte response body that follows it."""
    head, sep, rest = text.partition("\nResponse:")
    head = head.strip()
    # TestNG puts "expected [x] but found [y]" after the response body
    tail = rest.strip().splitlines()[-1].strip() if sep and rest.strip() else ""
    if tail.startswith("expected ["):
        head = f"{head} ({tail})"
    return head if len(head) <= limit else head[:limit] + "..."


def counts(rows):
    c = {k: sum(1 for r in rows if r["status"] == k) for k in ("Passed", "Failed", "Error", "Skipped")}
    c["Total"] = len(rows)
    return c


def run_meta():
    server, repo, run_id = (os.environ.get(k, "") for k in ("GITHUB_SERVER_URL", "GITHUB_REPOSITORY", "GITHUB_RUN_ID"))
    return {
        "url": f"{server}/{repo}/actions/runs/{run_id}" if run_id else "",
        "instance": os.environ.get("OHRM_BASE_URL") or "default instance (Credentials.java)",
        "when": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
    }


def write_sheet(rows, c, meta, path):
    wb = Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F28C38")
    header_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Summary"
    summary = [("Run date", meta["when"]), ("Instance", meta["instance"]), ("Run link", meta["url"]), ("", ""),
               ("Total", c["Total"]), ("Passed", c["Passed"]), ("Failed", c["Failed"]),
               ("Errors", c["Error"]), ("Skipped", c["Skipped"]),
               ("Pass rate", f"{c['Passed'] / c['Total'] * 100:.1f}%" if c["Total"] else "n/a")]
    for row in summary:
        ws.append(row)
        ws.cell(ws.max_row, 1).font = bold
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70

    ws = wb.create_sheet("Test Results")
    headers = ["TC ID", "Module", "Scenario", "Test Method", "Status", "Duration (s)", "Error Message"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    for r in rows:
        # Excel caps a cell at 32,767 characters
        ws.append([r["tc_id"], r["module"], r["scenario"], r["method"], r["status"], r["duration"], r["error"][:32000]])
        ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor=STATUS_FILL.get(r["status"], "FFFFFF"))
        ws.cell(ws.max_row, 7).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFG", (13, 18, 60, 40, 10, 12, 90)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_email(rows, c, meta, path):
    bad = [r for r in rows if r["status"] in ("Failed", "Error")]
    ok = not bad and c["Total"] > 0
    colour = "#2E7D32" if ok else "#C62828"
    e = html.escape

    cells = "".join(
        f'<td style="padding:8px 16px;border:1px solid #ddd;text-align:center"><b style="font-size:18px">{c[k]}</b><br>{label}</td>'
        for k, label in (("Total", "Total"), ("Passed", "Passed"), ("Failed", "Failed"), ("Error", "Errors"), ("Skipped", "Skipped")))
    parts = [
        '<div style="font-family:Arial,sans-serif;font-size:14px;color:#222">',
        f'<h2 style="color:{colour};margin-bottom:4px">OrangeHRM API Tests: {"PASSED" if ok else "FAILED"}</h2>',
        f'<p style="margin-top:0;color:#555">{e(meta["when"])} &middot; {e(meta["instance"])}</p>',
        f'<table style="border-collapse:collapse"><tr>{cells}</tr></table>',
    ]
    if bad:
        parts.append(f'<h3>Failed tests ({len(bad)})</h3>')
        parts.append('<table style="border-collapse:collapse;width:100%">'
                     '<tr style="background:#F28C38;color:#fff">'
                     + "".join(f'<th style="padding:6px;border:1px solid #ddd;text-align:left">{h}</th>'
                               for h in ("TC ID", "Module", "Scenario", "Error"))
                     + "</tr>")
        for r in bad:
            parts.append(
                "<tr>"
                f'<td style="padding:6px;border:1px solid #ddd;white-space:nowrap;vertical-align:top">{e(r["tc_id"] or r["method"])}</td>'
                f'<td style="padding:6px;border:1px solid #ddd;vertical-align:top">{e(r["module"])}</td>'
                f'<td style="padding:6px;border:1px solid #ddd;vertical-align:top">{e(r["scenario"])}</td>'
                f'<td style="padding:6px;border:1px solid #ddd;vertical-align:top;font-family:monospace;font-size:12px;white-space:pre-wrap">{e(first_lines(r["error"]))}</td>'
                "</tr>")
        parts.append("</table>")
    elif c["Total"] == 0:
        parts.append("<p><b>No test results were produced</b> - the build most likely failed before the tests ran. Check the run log.</p>")
    parts.append(f'<p>The full per-test results are in the attached <b>{SHEET_NAME}</b>; the PDF report is attached too.</p>')
    if meta["url"]:
        parts.append(f'<p><a href="{e(meta["url"])}">Open this run on GitHub</a></p>')
    parts.append("</div>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


def write_github_outputs(rows, c):
    bad = [r for r in rows if r["status"] in ("Failed", "Error")]
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        lines = ["## OrangeHRM API Test Results", "",
                 "| Total | Passed | Failed | Errors | Skipped |", "|---|---|---|---|---|",
                 f"| {c['Total']} | {c['Passed']} | {c['Failed']} | {c['Error']} | {c['Skipped']} |", ""]
        if bad:
            lines += ["### Failed tests", "", "| TC ID | Scenario | Error |", "|---|---|---|"]
            lines += [f"| {r['tc_id'] or r['method']} | {r['scenario']} | "
                      f"{first_lines(r['error'], 200).replace('|', '/').replace(chr(10), ' ')} |" for r in bad]
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as f:
            f.write(f"total={c['Total']}\npassed={c['Passed']}\nbad={len(bad)}\n")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    rows = load_results()
    c = counts(rows)
    meta = run_meta()
    write_sheet(rows, c, meta, os.path.join(OUT_DIR, SHEET_NAME))
    write_email(rows, c, meta, os.path.join(OUT_DIR, "email-body.html"))
    write_github_outputs(rows, c)
    print(f"Total {c['Total']} | Passed {c['Passed']} | Failed {c['Failed']} | Errors {c['Error']} | Skipped {c['Skipped']}")


if __name__ == "__main__":
    main()
